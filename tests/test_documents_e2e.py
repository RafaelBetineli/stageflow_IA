import tempfile
import unittest
from datetime import date
from pathlib import Path

from docx import Document

from stageflow.activities import ActivityAllocator, activity_document_fields
from stageflow.documents import MANUAL_CONTENT_FIELDS, DocumentService
from stageflow.template_engine import DocumentGenerator
from stageflow.extraction import RuleBasedExtractor
from stageflow.scheduling import HolidayProvider, ScheduleCalculator


PROJECT_ROOT = Path(__file__).resolve().parents[1]
EXAMPLE_TEXT = (PROJECT_ROOT / "data" / "mensagem_zap.example.txt").read_text(encoding="utf-8")


class DocumentServiceEndToEndTests(unittest.TestCase):
    def _data_for_area(self, area: str) -> dict[str, str]:
        text = EXAMPLE_TEXT.replace("Área estágio: Drogaria", f"Área estágio: {area}")
        if area == "Estética":
            text = text.replace("Conselho RT: CRF-SP 000000", "Conselho RT: CRBM-SP 000000")
            text = text.replace(
                "Cargo representante: Farmacêutico responsável técnico",
                "Cargo representante: Biomédica responsável técnica",
            )
        return RuleBasedExtractor().extract(text).fields

    def test_generates_three_complete_documents_for_every_area(self) -> None:
        for area in ("Estética", "Drogaria", "Hospitalar", "Manipulação", "Controle de qualidade"):
            with self.subTest(area=area), tempfile.TemporaryDirectory() as temporary:
                outputs = DocumentService().generate(self._data_for_area(area), temporary)

                self.assertEqual(3, len(outputs))
                for output in outputs:
                    document = Document(output)
                    self.assertFalse(DocumentGenerator._collect_placeholders(document))

    def test_manual_content_is_always_blank(self) -> None:
        data = self._data_for_area("Drogaria")
        data.update({field: "Texto que não deve ser usado" for field in MANUAL_CONTENT_FIELDS})
        with tempfile.TemporaryDirectory() as temporary:
            outputs = DocumentService().generate(data, temporary)

            report = Document(outputs[0])
            text = "\n".join(paragraph.text for paragraph in report.paragraphs)
            self.assertNotIn("Texto que não deve ser usado", text)

    def test_activity_titles_and_hours_are_objective_document_data(self) -> None:
        data = self._data_for_area("Drogaria")
        data.update(
            {
                "TITULO_ATV1": "Conferência de medicamentos",
                "CARGA_ATV1": "130 horas",
                "ATV1": "Narrativa que deve permanecer fora do documento",
            }
        )
        with tempfile.TemporaryDirectory() as temporary:
            outputs = DocumentService().generate(data, temporary)

            documents = tuple(Document(path) for path in outputs)
            xml_text = "\n".join(
                paragraph.text
                for document in documents
                for paragraph in document.paragraphs
            ) + "\n".join(
                cell.text
                for document in documents
                for table in document.tables
                for row in table.rows
                for cell in row.cells
            )
            self.assertIn("Conferência de medicamentos", xml_text)
            self.assertIn("130 horas", xml_text)
            self.assertNotIn("Narrativa que deve permanecer", xml_text)

    def test_report_frequency_table_is_filled_from_schedule(self) -> None:
        data = self._data_for_area("Drogaria")
        schedule = ScheduleCalculator().calculate(
            requested_start=date(2026, 9, 7),
            total_hours=130,
            excluded_dates=HolidayProvider().for_years((2026,), "SP"),
        )
        with tempfile.TemporaryDirectory() as temporary:
            outputs = DocumentService().generate(data, temporary, schedule=schedule)

            report = Document(outputs[0])
            frequency = report.tables[2]
            self.assertEqual("08/09/2026", frequency.cell(1, 1).text)
            self.assertEqual("08h00", frequency.cell(1, 2).text)
            self.assertEqual("14h00", frequency.cell(1, 3).text)
            self.assertEqual("6 horas", frequency.cell(1, 4).text)
            self.assertEqual("4 horas", frequency.cell(22, 4).text)
            self.assertEqual("", frequency.cell(1, 5).text)

    def test_report_uses_contextual_values_and_split_council(self) -> None:
        data = self._data_for_area("Drogaria")
        data.update(
            {
                "MODULO_ESTAGIO": "Módulo IV",
                "CAMPUS": "Campus Centro",
                "SEMESTRE": "6º semestre",
                "CONSELHO_RT": "CRF-SP 90909",
            }
        )
        with tempfile.TemporaryDirectory() as temporary:
            outputs = DocumentService().generate(data, temporary)

            report = Document(outputs[0])
            paragraphs = [paragraph.text for paragraph in report.paragraphs]
            self.assertIn("Módulo de estágio: IV", paragraphs)
            self.assertTrue(
                any(
                    text.startswith("SEMESTRE MATRICULADO")
                    and text.rstrip().endswith(": 6º")
                    for text in paragraphs
                )
            )
            self.assertIn("CAMPUS MATRICULADO: Centro", paragraphs)
            self.assertTrue(
                any(text.endswith("CRF") for text in paragraphs if text.startswith("SIGLA DO CONSELHO"))
            )
            self.assertIn("Nº DE INSCRIÇÃO NO CONSELHO: 90909", paragraphs)
            self.assertTrue(any("Conselho CRF nº 90909" in text for text in paragraphs))
            self.assertTrue(
                any(
                    "Registro Profissional (CRF, CRBM e etc) CRF nº 90909"
                    in text
                    for text in paragraphs
                )
            )

            plan = Document(outputs[1])
            self.assertIn("CRF nº 90909", [paragraph.text for paragraph in plan.paragraphs])

            insurer = next(
                paragraph
                for paragraph in report.paragraphs
                if paragraph.text.startswith("EMPRESA SEGURADORA:")
            )
            value_runs = [run for run in insurer.runs if "Seguradora Exemplo" in run.text]
            self.assertEqual(1, len(value_runs))
            self.assertFalse(value_runs[0].bold)

            term = Document(outputs[2])
            term_text = "\n".join(paragraph.text for paragraph in term.paragraphs)
            self.assertIn("no 6º semestre,", term_text)
            footer_xml = term.sections[0].footer._element.xml
            self.assertIn("NUMPAGES", footer_xml)

    def test_plan_fills_each_activity_hours_and_total(self) -> None:
        for area, total, count in (("Estética", 360, 10), ("Drogaria", 130, 10)):
            with self.subTest(area=area, total=total), tempfile.TemporaryDirectory() as temporary:
                data = self._data_for_area(area)
                allocations = ActivityAllocator().allocate(
                    total,
                    tuple((f"Atividade de exemplo {position}", position) for position in range(1, count + 1)),
                )
                data.update(activity_document_fields(allocations))
                data["CARGA_HORARIA"] = f"{total} horas"
                outputs = DocumentService().generate(data, temporary)

                plan = Document(outputs[1])
                activities = plan.tables[1]
                for position, allocation in enumerate(allocations, start=1):
                    self.assertEqual(allocation.title, activities.cell(position, 1).text)
                    self.assertEqual(f"{allocation.hours} horas", activities.cell(position, 2).text)
                for position in range(count + 1, len(activities.rows) - 1):
                    self.assertEqual("", activities.cell(position, 2).text)
                self.assertEqual(f"{total} horas", activities.rows[-1].cells[-1].text)

    def test_existing_document_is_preserved_without_overwrite(self) -> None:
        data = self._data_for_area("Drogaria")
        with tempfile.TemporaryDirectory() as temporary:
            output_dir = Path(temporary)
            existing = output_dir / "relatorio_aluno_exemplo.docx"
            existing.write_bytes(b"documento anterior")

            with self.assertRaisesRegex(FileExistsError, "Autorize a substituição"):
                DocumentService().generate(data, output_dir)

            self.assertEqual(b"documento anterior", existing.read_bytes())

    def test_overwrite_replaces_existing_documents(self) -> None:
        data = self._data_for_area("Drogaria")
        with tempfile.TemporaryDirectory() as temporary:
            output_dir = Path(temporary)
            expected = tuple(
                output_dir / f"{prefix}_aluno_exemplo.docx"
                for prefix in ("relatorio", "plano_estagio", "termo_compromisso")
            )
            for path in expected:
                path.write_bytes(b"documento anterior")

            outputs = DocumentService().generate(data, output_dir, overwrite=True)

            self.assertEqual(expected, outputs)
            for output in outputs:
                Document(output)


if __name__ == "__main__":
    unittest.main()
