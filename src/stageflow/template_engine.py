"""Substituição validada e atômica de placeholders em arquivos DOCX."""

from __future__ import annotations

from copy import deepcopy
from decimal import Decimal
import re
from pathlib import Path
from tempfile import NamedTemporaryFile
from typing import Protocol

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.text.run import Run
from docx.text.paragraph import Paragraph



class AttendanceDay(Protocol):
    day: object
    start: object
    end: object
    minutes: int


PLACEHOLDER_PATTERN = re.compile(r"\{\{\s*([^{}]+?)\s*\}\}")


class DocumentGenerationError(RuntimeError):
    """Erro de geração que impede a publicação de um DOCX incompleto."""


class MissingPlaceholderValueError(DocumentGenerationError):
    """O template exige chaves que não foram fornecidas."""


class MalformedPlaceholderError(DocumentGenerationError):
    """O template contém delimitadores de placeholder incompletos."""


class UnresolvedPlaceholderError(DocumentGenerationError):
    """O documento continuou contendo placeholders após o preenchimento."""


class DocumentGenerator:
    """Gera um DOCX somente quando todos os placeholders são resolvidos."""

    def __init__(self, template_path: str | Path):
        self.template_path = Path(template_path)

    def generate(
        self,
        data: dict,
        output_path: str | Path,
        *,
        attendance_days: tuple[AttendanceDay, ...] = (),
    ) -> None:
        output = Path(output_path)
        if not self.template_path.is_file():
            raise FileNotFoundError(f"Template não encontrado: {self.template_path}")

        doc = Document(self.template_path)
        malformed = self._collect_malformed_placeholders(doc)
        if malformed:
            raise MalformedPlaceholderError(
                "Marcadores malformados no template: " + " | ".join(malformed)
            )

        required = self._collect_placeholders(doc)
        missing = sorted(required - set(data))
        if missing:
            raise MissingPlaceholderValueError(
                "Valores ausentes para os placeholders: " + ", ".join(missing)
            )

        for paragraph in self._iter_paragraphs(doc):
            self._replace_in_paragraph(paragraph, data)

        self._fill_attendance_tables(doc, attendance_days)

        unresolved = sorted(self._collect_placeholders(doc))
        if unresolved:
            raise UnresolvedPlaceholderError(
                "Placeholders não resolvidos: " + ", ".join(unresolved)
            )

        output.parent.mkdir(parents=True, exist_ok=True)
        temporary_path: Path | None = None
        try:
            with NamedTemporaryFile(
                suffix=".docx",
                prefix=f".{output.stem}_",
                dir=output.parent,
                delete=False,
            ) as temporary:
                temporary_path = Path(temporary.name)
            doc.save(temporary_path)
            Document(temporary_path)
            temporary_path.replace(output)
        except Exception as error:
            if temporary_path is not None:
                temporary_path.unlink(missing_ok=True)
            if isinstance(error, DocumentGenerationError):
                raise
            raise DocumentGenerationError(
                f"Não foi possível gerar {output.name}: {error}"
            ) from error

    @classmethod
    def _fill_attendance_tables(
        cls,
        doc,
        days: tuple[AttendanceDay, ...],
    ) -> None:
        if not days:
            return
        for table in doc.tables:
            for row in table.rows:
                cells = row.cells
                if len(cells) < 5 or not cells[0].text.strip().isdigit():
                    continue
                position = int(cells[0].text.strip()) - 1
                if position < 0 or position >= len(days):
                    continue
                item = days[position]
                hours = cls._format_minutes(item.minutes)
                unit = "hora" if hours == "1" else "horas"
                values = (
                    item.day.strftime("%d/%m/%Y"),
                    item.start.strftime("%Hh%M"),
                    item.end.strftime("%Hh%M"),
                    f"{hours} {unit}",
                )
                for cell, value in zip(cells[1:5], values):
                    cls._set_cell_value(cell, value)

    @staticmethod
    def _set_cell_value(cell, value: str) -> None:
        paragraph = cell.paragraphs[0]
        if paragraph.runs:
            paragraph.runs[0].text = value
            for run in paragraph.runs[1:]:
                run.text = ""
        else:
            paragraph.add_run(value)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    @staticmethod
    def _format_minutes(minutes: int) -> str:
        value = (Decimal(minutes) / Decimal(60)).quantize(Decimal("0.01"))
        return format(value, "f").rstrip("0").rstrip(".")

    @classmethod
    def _collect_placeholders(cls, doc) -> set[str]:
        return {
            match.group(1).strip()
            for paragraph in cls._iter_paragraphs(doc)
            for match in PLACEHOLDER_PATTERN.finditer(paragraph.text)
        }

    @classmethod
    def _collect_malformed_placeholders(cls, doc) -> tuple[str, ...]:
        malformed: list[str] = []
        for paragraph in cls._iter_paragraphs(doc):
            residual = PLACEHOLDER_PATTERN.sub("", paragraph.text)
            if "{{" in residual or "}}" in residual:
                malformed.append(paragraph.text.strip())
        return tuple(dict.fromkeys(malformed))

    @classmethod
    def _iter_paragraphs(cls, doc):
        yield from doc.paragraphs
        yield from cls._iter_table_paragraphs(doc.tables)
        yield from cls._iter_text_box_paragraphs(doc.element.body, doc)
        for section in doc.sections:
            for container in (section.header, section.footer):
                yield from container.paragraphs
                yield from cls._iter_table_paragraphs(container.tables)
                yield from cls._iter_text_box_paragraphs(
                    container._element,
                    container,
                )

    @staticmethod
    def _iter_text_box_paragraphs(element, parent):
        for paragraph_xml in element.xpath(".//w:txbxContent//w:p"):
            yield Paragraph(paragraph_xml, parent)

    @classmethod
    def _iter_table_paragraphs(cls, tables):
        seen_cells: set[object] = set()
        for table in tables:
            for row in table.rows:
                for cell in row.cells:
                    cell_xml = cell._tc
                    if cell_xml in seen_cells:
                        continue
                    seen_cells.add(cell_xml)
                    yield from cell.paragraphs
                    yield from cls._iter_table_paragraphs(cell.tables)

    def _replace_in_paragraph(self, paragraph, data: dict) -> None:
        runs = self._paragraph_runs(paragraph)
        full_text = "".join(run.text for run in runs)
        keys = {
            match.group(1).strip()
            for match in PLACEHOLDER_PATTERN.finditer(full_text)
        }
        for key in keys:
            placeholder = re.compile(r"\{\{\s*" + re.escape(key) + r"\s*\}\}")
            while True:
                runs = self._paragraph_runs(paragraph)
                full_text = "".join(run.text for run in runs)
                match = placeholder.search(full_text)
                if match is None:
                    break
                value = str(data[key])
                blocks = self._split_blocks(value)
                if len(blocks) > 1 and PLACEHOLDER_PATTERN.fullmatch(
                    full_text.strip()
                ):
                    self._replace_with_paragraphs(
                        paragraph,
                        match.start(),
                        match.end(),
                        blocks,
                    )
                    break
                self._replace_range(
                    paragraph,
                    match.start(),
                    match.end(),
                    value,
                )

    @staticmethod
    def _split_blocks(value: str) -> tuple[str, ...]:
        if "\n" not in value and "\r" not in value:
            return (value,)
        return tuple(
            block.strip()
            for block in re.split(r"(?:\r?\n)+", value)
            if block.strip()
        ) or ("",)

    @classmethod
    def _replace_with_paragraphs(
        cls,
        paragraph,
        start: int,
        end: int,
        blocks: tuple[str, ...],
    ) -> None:
        template_xml = deepcopy(paragraph._p)
        cls._replace_range(paragraph, start, end, blocks[0])

        current_xml = paragraph._p
        for block in blocks[1:]:
            cloned_xml = deepcopy(template_xml)
            current_xml.addnext(cloned_xml)
            cloned = Paragraph(cloned_xml, paragraph._parent)
            placeholder = PLACEHOLDER_PATTERN.search(cloned.text)
            if placeholder is None:
                raise DocumentGenerationError(
                    "Não foi possível clonar o parágrafo do placeholder"
                )
            cls._replace_range(
                cloned,
                placeholder.start(),
                placeholder.end(),
                block,
            )
            current_xml = cloned_xml

    @staticmethod
    def _paragraph_runs(paragraph) -> list[Run]:
        """Inclui runs dentro de hyperlinks e controles pertencentes ao parágrafo."""
        runs: list[Run] = []
        for element in paragraph._p.xpath(".//w:r"):
            nearest_paragraph = next(element.iterancestors(qn("w:p")), None)
            if nearest_paragraph is paragraph._p:
                runs.append(Run(element, paragraph))
        return runs

    @classmethod
    def _replace_range(cls, paragraph, start: int, end: int, value: str) -> None:
        runs = cls._paragraph_runs(paragraph)
        offset = 0
        first_index = None
        last_index = None
        prefix = ""
        suffix = ""

        for index, run in enumerate(runs):
            run_start = offset
            run_end = offset + len(run.text)
            if first_index is None and run_end > start:
                first_index = index
                prefix = run.text[: start - run_start]
            if first_index is not None and run_end >= end:
                last_index = index
                suffix = run.text[end - run_start :]
                break
            offset = run_end

        if first_index is None or last_index is None:
            raise DocumentGenerationError("Não foi possível mapear um placeholder nos runs")

        runs[first_index].text = prefix + value + suffix
        for index in range(first_index + 1, last_index + 1):
            runs[index].text = ""
